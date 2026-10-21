"""Load a Clover merge review sheet into the in-app review, where decisions stay editable.

Reads the `review` sheet of an .xlsx (or a .csv) with at least legacy_id, jizhanming, legacy_name
and decision columns; filled_by, confidence and why are kept when present. A row is loaded only
when its legacy id still names the same product (记账名 and name) in this database. Rows already
edited or applied in the app are left alone unless --replace (applied rows are never replaced).
Nothing is merged; managers apply decisions from Products → Clover matching.

    python scripts/load_clover_match_review.py --sheet review.xlsx [--db PATH] [--dry-run] [--replace]
"""
import argparse
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'popcore_app'))


def read_sheet(path):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with open(path, newline='', encoding='utf-8-sig') as handle:
            return list(csv.DictReader(handle))
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book['review'] if 'review' in book.sheetnames else book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else '' for cell in next(rows)]
    if 'legacy_id' not in header or 'decision' not in header:
        raise ValueError('The sheet needs legacy_id and decision columns')
    return [dict(zip(header, values)) for values in rows if any(v not in (None, '') for v in values)]


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    parser.add_argument('--sheet', required=True)
    parser.add_argument('--db')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--replace', action='store_true', help='also overwrite decisions edited in the app')
    args = parser.parse_args(argv)
    import db  # noqa: E402
    if args.db:
        db.DB_PATH = str(Path(args.db).resolve())
    db.migrate_db()
    from clover_match_review import load_rows  # noqa: E402
    import merge_clover_products as merger  # noqa: E402
    rows = read_sheet(args.sheet)
    con = merger.connect(db.DB_PATH)
    try:
        outcomes = load_rows(con, rows, 'sheet:' + Path(args.sheet).name, replace=args.replace)
        if args.dry_run:
            con.rollback()
        else:
            con.commit()
    finally:
        con.close()
    print(f"{len(rows)} sheet rows{' (dry run, nothing saved)' if args.dry_run else ''}: "
          + ', '.join(f'{k} {v}' for k, v in sorted(outcomes.items())))
    return 0


if __name__ == '__main__':
    sys.exit(main())
